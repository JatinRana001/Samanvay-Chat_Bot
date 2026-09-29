"""Import the Maharashtra approvals workbook into the structured registry."""
import argparse
import csv
import json
import os
import re
import sys
import uuid
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app.database import SessionLocal  # noqa: E402
from app.models import (  # noqa: E402
    Approval, ApprovalDocument, ApprovalRule, Department, Document, Industry, Location, RegulatoryDocument, RegulatoryChunk,
)
from app.industry_taxonomy import CANONICAL_INDUSTRIES, normalize_industry, normalize_rule_industry  # noqa: E402
from app.rag.chunker import TextChunker  # noqa: E402
from app.rag.embedder import EmbeddingService  # noqa: E402

DEFAULT_WORKBOOK = Path(__file__).resolve().parents[1] / "data" / "Maharashtra_Industry_Approvals_Database.xlsx"
EXPECTED_COUNTS = {"Industries": 20, "Approvals": 20, "Approval_Rules": 53, "Documents": 23, "Locations": 30}


def sheet_rows(workbook, sheet_name):
    ws = workbook[sheet_name]
    rows = list(ws.values)
    header_index = next((i for i, row in enumerate(rows) if any(v is not None for v in row)), None)
    if header_index is None:
        return []
    headers = [str(v).strip() if v is not None else "" for v in rows[header_index]]
    return [dict(zip(headers, row), __source_row__=i + 1) for i, row in enumerate(rows[header_index + 1:], start=header_index + 1)
            if any(v is not None for v in row)]


def val(row, *keys, default=None):
    normalized = {str(k).strip().lower().replace(" ", "_"): v for k, v in row.items() if k}
    for key in keys:
        value = normalized.get(key.lower().replace(" ", "_"))
        if value is not None and str(value).strip() != "":
            return value
    return default


def text(value):
    return str(value).strip() if value is not None and str(value).strip() else None


def boolean(value, default=None):
    if value is None or str(value).strip() == "":
        return default
    normalized = str(value).strip().casefold()
    if normalized in {"yes", "y", "true", "1"} or normalized.startswith("yes "):
        return True
    if normalized in {"no", "n", "false", "0"} or normalized.startswith("no "):
        return False
    return default


def upsert(db, model, row, cache, reset=False):
    if model not in cache:
        cache[model] = {} if reset else {obj.id: obj for obj in db.query(model).all()}
    existing = cache[model].get(row["id"])
    if existing is None:
        existing = model(**row)
        db.add(existing)
        cache[model][row["id"]] = existing
    else:
        for key, value in row.items():
            setattr(existing, key, value)


def import_industries(db, wb, cache, reset=False):
    for r in sheet_rows(wb, "Industries"):
        ident = text(val(r, "industry_id"))
        if ident:
            raw_name = text(val(r, "industry_name")) or ident
            industry_name, _ = normalize_industry(raw_name)
            if industry_name not in CANONICAL_INDUSTRIES:
                raise ValueError(f"Industry row {ident!r} uses unknown industry name {raw_name!r}")
            summary = text(val(r, "basic_requirements"))
            upsert(db, Industry, {"id": ident, "name": industry_name,
                                  "sector": text(val(r, "sector")) or "Unspecified",
                                  "description": text(val(r, "description")),
                                  "basic_setup_information": {"summary": summary} if summary else {}},
                  cache, reset)


def import_locations(db, wb, cache, reset=False):
    for r in sheet_rows(wb, "Locations"):
        ident = text(val(r, "location_id"))
        if ident:
            zone = text(val(r, "special_zone"))
            upsert(db, Location, {"id": ident, "state": "Maharashtra",
                                  "district": text(val(r, "district")) or "Unknown",
                                  "taluka": text(val(r, "taluka")), "city": text(val(r, "city")),
                                  "industrial_area": text(val(r, "industrial_area")),
                                  "zone": zone,
                                  "special_conditions": {"summary": zone} if zone else {}},
                  cache, reset)


def import_documents(db, wb, cache, reset=False):
    for r in sheet_rows(wb, "Documents"):
        ident = text(val(r, "document_id", "doc_id"))
        if ident:
            upsert(db, Document, {"id": ident, "name": text(val(r, "document_name", "name")) or ident,
                                  "description": text(val(r, "description")),
                                  "issuing_authority": text(val(r, "issuing_authority", "who_issues_it")),
                                  "validity": text(val(r, "validity")), "format": text(val(r, "format")),
                                  "notes": text(val(r, "notes")), "is_demo": False,
                                  "last_verified": None, "source_workbook": DEFAULT_WORKBOOK.name,
                                  "source_row": r.get("__source_row__"), "imported_at": datetime.utcnow()}, cache, reset)


def import_approvals(db, wb, cache, reset=False):
    departments = {} if reset else {dept.name: dept.id for dept in db.query(Department).all()}
    for r in sheet_rows(wb, "Approvals"):
        ident = text(val(r, "approval_id"))
        if not ident:
            continue
        department_name = text(val(r, "department"))
        dept_id = None
        if department_name:
            dept_id = departments.get(department_name)
            if dept_id is None:
                dept_id = str(uuid.uuid4())
                db.add(Department(id=dept_id, name=department_name))
                departments[department_name] = dept_id
        verified = val(r, "last_verified")
        if isinstance(verified, datetime):
            verified = verified.date()
        if not isinstance(verified, date):
            verified = None
        description = text(val(r, "description"))
        process = text(val(r, "application_process"))
        renewal_notes = text(val(r, "renewal_notes"))
        if process:
            description = f"{description or ''}\nApplication process: {process}".strip()
        if renewal_notes:
            description = f"{description or ''}\nRenewal notes: {renewal_notes}".strip()
        portal_raw = text(val(r, "official_portal", "portal"))
        portal_urls = []
        for part in re.split(r"\s+/\s+|\s+via\s+", portal_raw or "", flags=re.IGNORECASE):
            candidate = part.strip().rstrip(".,;")
            if candidate and "." in candidate:
                portal_urls.append(candidate if re.match(r"https?://", candidate, re.I) else f"https://{candidate}")
        upsert(db, Approval, {"id": ident, "name": text(val(r, "approval_name", "name")) or ident,
                              "description": description, "purpose": text(val(r, "purpose")),
                              "department_id": dept_id,
                              "official_portal": portal_urls[0] if portal_urls else None,
                              "official_portal_urls": portal_urls, "official_portal_raw": portal_raw,
                              "official_source": text(val(r, "official_source", "source")) or "Source workbook",
                              "fee": text(val(r, "fee")), "processing_time": text(val(r, "processing_time")),
                              "validity": text(val(r, "validity")),
                              "renewal_required": boolean(val(r, "renewal_required"), default=False),
                              "last_verified": verified, "source_workbook": DEFAULT_WORKBOOK.name,
                              "source_row": r.get("__source_row__"), "imported_at": datetime.utcnow(),
                              "status": text(val(r, "status")) or "ACTIVE",
                              "is_demo": False}, cache, reset)
    db.flush()
    return departments


def import_approval_rules(db, wb, cache, reset=False):
    for r in sheet_rows(wb, "Approval_Rules"):
        ident = text(val(r, "rule_id"))
        approval_id = text(val(r, "approval_id"))
        if not ident:
            continue
        if not approval_id or approval_id not in cache.get(Approval, {}):
            raise ValueError(f"Rule row {ident!r} references unknown approval_id {approval_id!r}")
        raw_industry = text(val(r, "industry"))
        sub_sector = text(val(r, "sub_sector"))
        if raw_industry and raw_industry.lower().startswith("all "):
            industry = None
        else:
            industry, sub_sector = normalize_rule_industry(raw_industry, sub_sector)
            if raw_industry and industry not in CANONICAL_INDUSTRIES:
                raise ValueError(
                    f"Approval rule {ident!r} uses unknown industry label {raw_industry!r}; "
                    "add an explicit canonical alias before importing."
                )
        district = text(val(r, "district"))
        if industry and industry.lower().startswith("all "):
            industry = None
        if district and district.lower().startswith("all "):
            district = None
        applicability = text(val(r, "applicability")) or "Information required"
        upsert(db, ApprovalRule, {"id": ident, "approval_id": approval_id, "industry": industry,
                                  "industry_label_raw": raw_industry,
                                  "sub_sector": sub_sector, "activity": text(val(r, "activity")),
                                  "state": text(val(r, "state")) or "Maharashtra", "district": district,
                                  "investment_min": val(r, "investment_min"),
                                  "investment_max": val(r, "investment_max"),
                                  "employee_min": val(r, "employee_min", "employees_min"),
                                  "employee_max": val(r, "employee_max", "employees_max"),
                                  "project_stage": text(val(r, "project_stage")),
                                  "construction_required": boolean(val(r, "construction_required")),
                                  "hazardous_materials": boolean(val(r, "hazardous_materials")),
                                  "pollution_category": text(val(r, "pollution_category")),
                                  "conditions": _json_value(val(r, "conditions"), ident),
                                  "applicability": applicability, "is_demo": False}, cache, reset)


def import_approval_documents(db, cache):
    """Import explicitly grounded indicative document mappings from CSV."""
    db.query(ApprovalDocument).delete()
    path = Path(__file__).resolve().parents[1] / "data" / "approval_documents.csv"
    with path.open(encoding="utf-8-sig", newline="") as stream:
        for row in csv.DictReader(stream):
            approval_id, document_id = row["approval_id"].strip(), row["document_id"].strip()
            if approval_id not in cache.get(Approval, {}):
                raise ValueError(f"Document mapping references unknown approval_id {approval_id!r}")
            if document_id not in cache.get(Document, {}):
                raise ValueError(f"Document mapping references unknown document_id {document_id!r}")
            existing = db.query(ApprovalDocument).filter_by(approval_id=approval_id, document_id=document_id).first()
            values = {"mandatory": boolean(row["mandatory"], False), "condition": row["condition"] or None,
                      "source": row["source"], "needs_review": boolean(row["needs_review"], True)}
            if existing:
                for key, value in values.items():
                    setattr(existing, key, value)
            else:
                db.add(ApprovalDocument(approval_id=approval_id, document_id=document_id, **values))


def import_workbook_rag(db, cache):
    """Index workbook text for retrieval fallback using a fixed local hash space."""
    records = []
    for approval in cache.get(Approval, {}).values():
        body = "\n".join(filter(None, [approval.name, approval.description, approval.purpose, approval.fee,
                                       approval.processing_time, approval.validity, approval.official_source]))
        records.append((f"WB-APR-{approval.id}", f"Approval: {approval.name}", "Approvals", body))
    for industry in cache.get(Industry, {}).values():
        body = "\n".join(filter(None, [industry.name, industry.sector, industry.description,
                                       (industry.basic_setup_information or {}).get("summary")]))
        records.append((f"WB-IND-{industry.id}", f"Industry: {industry.name}", "Industries", body))
    approval_names = {item.id: item.name for item in cache.get(Approval, {}).values()}
    for rule in cache.get(ApprovalRule, {}).values():
        notes = (rule.conditions or {}).get("notes") if isinstance(rule.conditions, dict) else None
        body = "\n".join(filter(None, [approval_names.get(rule.approval_id), rule.industry_label_raw,
                                       rule.sub_sector, rule.activity, rule.district, rule.project_stage,
                                       rule.pollution_category, rule.applicability, notes]))
        records.append((f"WB-RUL-{rule.id}", f"Rule {rule.id}: {approval_names.get(rule.approval_id, '')}", "Approval_Rules", body))

    for ident, title, sheet, body in records:
        if not body:
            continue
        document = db.query(RegulatoryDocument).filter_by(id=ident).first()
        if document is None:
            document = RegulatoryDocument(id=ident, title=title, department=None,
                source_url=f"workbook://{sheet}/{ident}", last_verified=None, extracted_text=body,
                category=sheet, state="Maharashtra", is_demo=False)
            db.add(document)
            db.flush()
        else:
            document.title, document.extracted_text = title, body
            db.query(RegulatoryChunk).filter_by(regulatory_document_id=ident).delete()
        for index, content in enumerate(TextChunker.chunk_text(body)):
            vector = EmbeddingService._fallback_embedding(content, EmbeddingService.DIMENSION)
            db.add(RegulatoryChunk(regulatory_document_id=ident, chunk_index=index, chunk_text=content,
                embedding_vector=vector, embedding_model="fallback-hash", embedding_dim=len(vector),
                metadata_json={"embedding_model": "fallback-hash", "embedding_dim": len(vector), "source_sheet": sheet}))


def _json_value(value, row_id="unknown"):
    if isinstance(value, dict):
        return value
    if not value:
        return {}
    try:
        parsed = json.loads(value)
        if not isinstance(parsed, dict):
            raise ValueError(f"Rule {row_id!r} conditions must be a JSON object; got {type(parsed).__name__}")
        return parsed
    except (TypeError, json.JSONDecodeError) as exc:
        raise ValueError(f"Rule {row_id!r} contains invalid conditions JSON: {value!r}") from exc


def normalize_existing_rule_industries(db):
    """Migrate legacy rule labels in place while retaining qualifier scope."""
    changed = 0
    for rule in db.query(ApprovalRule).all():
        if not rule.industry:
            continue
        canonical, sub_sector = normalize_rule_industry(rule.industry, rule.sub_sector)
        if canonical not in CANONICAL_INDUSTRIES:
            raise ValueError(
                f"Approval rule {rule.id!r} uses unknown industry label {rule.industry!r}; "
                "add an explicit canonical alias before importing."
            )
        if rule.industry != canonical or rule.sub_sector != sub_sector:
            rule.industry = canonical
            rule.sub_sector = sub_sector
            changed += 1
    return changed


def import_workbook(path, reset=False):
    wb = load_workbook(path, read_only=True, data_only=True)
    required = {"Industries", "Approvals", "Approval_Rules", "Documents", "Locations"}
    missing = required - set(wb.sheetnames)
    if missing:
        raise ValueError(f"Workbook is missing sheets: {', '.join(sorted(missing))}")
    expected_counts = {"Industries": 20, "Approvals": 20, "Approval_Rules": 53, "Documents": 23, "Locations": 30}
    actual_counts = {sheet: len(sheet_rows(wb, sheet)) for sheet in expected_counts}
    if actual_counts != expected_counts:
        raise ValueError(f"Workbook row counts are {actual_counts}; expected {expected_counts}")
    db = SessionLocal()
    try:
        if reset:
            for model in (RegulatoryChunk, RegulatoryDocument, ApprovalDocument, ApprovalRule, Approval, Department, Document, Industry, Location):
                db.execute(delete(model))
            db.flush()
        cache = {}
        import_industries(db, wb, cache, reset)
        import_locations(db, wb, cache, reset)
        import_documents(db, wb, cache, reset)
        import_approvals(db, wb, cache, reset)
        import_approval_rules(db, wb, cache, reset)
        import_approval_documents(db, cache)
        import_workbook_rag(db, cache)
        normalized_rules = normalize_existing_rule_industries(db)
        db.flush()
        imported_counts = {
            "Industries": db.query(Industry).count(),
            "Approvals": db.query(Approval).filter(Approval.is_demo.is_(False)).count(),
            "Approval_Rules": db.query(ApprovalRule).count(),
            "Documents": db.query(Document).filter(Document.is_demo.is_(False)).count(),
            "Locations": db.query(Location).count(),
        }
        if imported_counts != expected_counts:
            raise ValueError(f"Imported row counts are {imported_counts}; expected {expected_counts}. Use --reset for a clean import.")
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()
        wb.close()
    print("Imported counts: " + ", ".join(f"{key}={value}" for key, value in imported_counts.items()))
    print(f"Canonicalized legacy approval rule industry labels: {normalized_rules}")
    print("Maharashtra workbook import completed.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", default=os.getenv("MAHARASHTRA_WORKBOOK", str(DEFAULT_WORKBOOK)))
    parser.add_argument("--reset", action="store_true", help="Clear rows from these registry tables before importing")
    args = parser.parse_args()
    import_workbook(args.workbook, reset=args.reset)
